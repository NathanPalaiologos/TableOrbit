"""Small local format fixtures. No private inputs or downloads."""
import importlib.util
from pathlib import Path
import tempfile
import unittest
import datetime
import json
import subprocess
import os
import sys
import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.feather as feather

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('read_table',ROOT/'scripts/read_table.py')
reader=importlib.util.module_from_spec(spec);spec.loader.exec_module(reader)

class FileFormats(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='tableorbit-formats-')
        self.root=Path(self.temp.name).resolve()
    def tearDown(self):
        assert self.root==Path(self.temp.name).resolve()
        self.temp.cleanup()
    def test_excel_worksheet_selection_and_cached_values(self):
        file=self.root/'book.xlsx';book=openpyxl.Workbook();sheet=book.active;sheet.title='Events'
        sheet.append(['id','date','flag','value']);sheet.append(['001',datetime.datetime(2024,1,1),False,0])
        second=book.create_sheet('Notes');second.append(['name','note']);second.append(['A','中文'])
        book.create_sheet('Empty')
        book.save(file);book.close()
        result=reader.read_table(file)
        self.assertEqual(result['sheets'],['Events','Notes','Empty']);self.assertEqual(result['table']['rows'][0],{'id':'001','date':'2024-01-01T00:00:00','flag':'false','value':'0'})
        self.assertEqual(reader.read_table(file,'Notes')['table']['rows'][0]['note'],'中文')
        self.assertEqual(reader.read_table(file,'Empty')['table'],{'headers':[],'rows':[]})
        with self.assertRaisesRegex(ValueError,'Worksheet'):reader.read_table(file,'missing')
    def test_arrow_and_stata_preserve_record_count_and_identifiers(self):
        file=self.root/'records.feather'
        feather.write_feather(pa.table({'id':pa.array([9007199254740993,9007199254740994],pa.int64()),'flag':[False,True]}),file)
        result=reader.read_table(file)['table'];self.assertEqual(len(result['rows']),2)
        self.assertEqual(result['rows'][0]['id'],'9007199254740993');self.assertEqual(result['rows'][0]['flag'],'false')
        file=self.root/'records.dta';pd.DataFrame({'id':['001','001'],'value':[0.,2.]}).to_stata(file,write_index=False,version=118)
        result=reader.read_table(file)['table'];self.assertEqual(len(result['rows']),2)
        self.assertEqual(result['rows'][0]['id'],'001');self.assertEqual(float(result['rows'][0]['value']),0)
    def test_limits_fail_without_returning_a_partial_table(self):
        file=self.root/'large.feather';feather.write_feather(pa.table({'id':range(10)}),file)
        previous=reader.MAX_ROWS;reader.MAX_ROWS=5
        try:
            with self.assertRaisesRegex(ValueError,'limit'):reader.read_table(file)
        finally:reader.MAX_ROWS=previous
    def test_node_open_routes_formats_and_reports_errors(self):
        csv=self.root/'sample.csv';csv.write_text('id;note\n001;中文',encoding='utf-16')
        code="const {readData}=require('./scripts/files.cjs'); readData(process.argv[1]).then(x=>console.log(JSON.stringify(x))).catch(e=>{console.error(e.message);process.exitCode=1});"
        result=subprocess.run(['node','-e',code,str(csv)],cwd=ROOT,capture_output=True,text=True,encoding='utf-8',check=True)
        self.assertIn('中文',json.loads(result.stdout)['text'])
        arrow=self.root/'sample.arrow'
        with pa.OSFile(str(arrow),'wb') as stream:
            with pa.ipc.new_stream(stream,pa.schema([('id',pa.string())])) as writer:writer.write_table(pa.table({'id':['001']}))
        result=subprocess.run(['node','-e',code,str(arrow)],cwd=ROOT,env={**os.environ,'PANEL_SPACE_PYTHON':sys.executable},capture_output=True,text=True,encoding='utf-8',check=True)
        self.assertEqual(json.loads(result.stdout)['table']['rows'][0]['id'],'001')
        bad=self.root/'broken.xlsx';bad.write_text('not a spreadsheet')
        result=subprocess.run(['node','-e',code,str(bad)],cwd=ROOT,env={**os.environ,'PANEL_SPACE_PYTHON':sys.executable},capture_output=True,text=True,encoding='utf-8')
        self.assertNotEqual(result.returncode,0);self.assertTrue(result.stderr.strip())

if __name__=='__main__':unittest.main()
